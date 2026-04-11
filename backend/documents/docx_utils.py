"""
Utilities for DOCX template manipulation.
Supports {{placeholder}} syntax.
"""
import re
import os
import shutil
import tempfile

PLACEHOLDER_RE = re.compile(r'\{\{(\w+)\}\}')


def _iter_runs(doc):
    """Yield all runs from paragraphs and table cells."""
    for para in doc.paragraphs:
        yield from para.runs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    yield from para.runs


def extract_placeholders(file_path: str) -> list[str]:
    """
    Return sorted list of unique placeholder names found in a DOCX file.
    Returns empty list for non-DOCX files.
    """
    if not file_path.lower().endswith('.docx'):
        return []
    try:
        from docx import Document as DocxDocument
        doc = DocxDocument(file_path)
        found = set()
        for run in _iter_runs(doc):
            for m in PLACEHOLDER_RE.finditer(run.text):
                found.add(m.group(1))
        # Also scan full paragraph text (handles split runs)
        for para in doc.paragraphs:
            for m in PLACEHOLDER_RE.finditer(para.text):
                found.add(m.group(1))
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        for m in PLACEHOLDER_RE.finditer(para.text):
                            found.add(m.group(1))
        return sorted(found)
    except Exception:
        return []


def _replace_in_paragraph(para, fields: dict):
    """
    Replace placeholders in a paragraph, handling the case where
    a placeholder is split across multiple runs.
    """
    full_text = para.text
    if not PLACEHOLDER_RE.search(full_text):
        return

    # Build new full text with replacements
    def replacer(m):
        return str(fields.get(m.group(1), m.group(0)))

    new_text = PLACEHOLDER_RE.sub(replacer, full_text)

    # Put all text in first run, clear the rest
    if para.runs:
        para.runs[0].text = new_text
        for run in para.runs[1:]:
            run.text = ''


def fill_placeholders(src_path: str, fields: dict, dest_path: str):
    """
    Copy src DOCX, replace all {{key}} with fields[key], save to dest_path.
    Non-DOCX files are copied as-is.
    """
    if not src_path.lower().endswith('.docx'):
        shutil.copy2(src_path, dest_path)
        return

    from docx import Document as DocxDocument
    doc = DocxDocument(src_path)

    for para in doc.paragraphs:
        _replace_in_paragraph(para, fields)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    _replace_in_paragraph(para, fields)

    doc.save(dest_path)
